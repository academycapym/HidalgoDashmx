#!/usr/bin/env python3
"""Genera inventario_occidente.json desde Inventario_OCCIDENTE.xlsx."""
import json
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("Instala openpyxl: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "Inventario_OCCIDENTE.xlsx"
OUT = ROOT / "inventario_occidente.json"


def clean_id(val):
    s = str(val).strip().split(".")[0].lstrip("0")
    return s if s else "0"


def main():
    if not XLSX.exists():
        print(f"No existe {XLSX}", file=sys.stderr)
        sys.exit(1)

    wb = load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    headers = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        if v:
            headers[c] = str(v).strip()

    stock_col = max(c for c, h in headers.items() if h.endswith("_STOCK"))
    id_col = next(c for c, h in headers.items() if h.upper() == "IDPDV")
    mod_col = next(c for c, h in headers.items() if h.upper() == "DESCRIPCION")
    status_col = next(c for c, h in headers.items() if h.upper() == "STATUS_SKU")
    mar_col = next(c for c, h in headers.items() if h.upper() == "MARCA")
    stock_label = headers[stock_col]

    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        sid = clean_id(row[id_col - 1])
        if not sid.isdigit() or len(sid) < 5:
            continue
        try:
            q = int(float(row[stock_col - 1] or 0))
        except (TypeError, ValueError):
            q = 0
        rows.append({
            "id": sid,
            "mod": str(row[mod_col - 1] or "").strip(),
            "mar": str(row[mar_col - 1] or "").strip(),
            "status": str(row[status_col - 1] or "").strip().upper(),
            "q": q,
        })

    payload = {"actualizado_al": stock_label, "list": rows}
    OUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"OK: {len(rows)} registros -> {OUT} ({OUT.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
